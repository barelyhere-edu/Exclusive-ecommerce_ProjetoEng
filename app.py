"""
Backend em Python (Flask) para um site de compras.
Usa uma planilha Excel (data/produtos.xlsx) como "banco de dados":
- Le os produtos da planilha e envia para o front-end em JSON.
- Permite adicionar, editar e remover produtos, salvando direto na planilha.
- Tem uma rota para o front-end baixar a planilha atualizada.

Como rodar:
    pip install -r requirements.txt
    python app.py

A API vai rodar em http://localhost:5000
"""

import os
import threading
from flask import Flask, jsonify, request, send_file, send_from_directory, abort
from flask_cors import CORS
from openpyxl import Workbook, load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "data", "produtos.xlsx")
SHEET_NAME = "produtos"
COLUNAS = ["id", "nome", "categoria", "preco", "estoque", "descricao", "imagem_url"]

app = Flask(__name__)
CORS(app)  # permite que o front-end (rodando em outra porta/dominio) acesse a API

# Trava para evitar que duas requisicoes escrevam na planilha ao mesmo tempo
lock = threading.Lock()


# ---------------------------------------------------------------------------
# Funcoes auxiliares para ler e escrever no Excel
# ---------------------------------------------------------------------------

def garantir_planilha():
    """Cria a planilha com o cabecalho caso ela nao exista ainda."""
    if not os.path.exists(EXCEL_PATH):
        os.makedirs(os.path.dirname(EXCEL_PATH), exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(COLUNAS)
        wb.save(EXCEL_PATH)


def ler_produtos():
    """Le todas as linhas da planilha e devolve uma lista de dicionarios."""
    garantir_planilha()
    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]

    produtos = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:  # linha vazia
            continue
        produto = dict(zip(COLUNAS, row))
        produtos.append(produto)
    return produtos


def salvar_produtos(produtos):
    """Reescreve a planilha inteira a partir de uma lista de dicionarios."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(COLUNAS)
    for p in produtos:
        ws.append([p.get(col) for col in COLUNAS])
    wb.save(EXCEL_PATH)


def proximo_id(produtos):
    if not produtos:
        return 1
    return max(p["id"] for p in produtos) + 1


# ---------------------------------------------------------------------------
# Rotas da API (consumidas pelo front-end)
# ---------------------------------------------------------------------------

@app.route("/api/produtos", methods=["GET"])
def listar_produtos():
    """Lista todos os produtos. Aceita filtros opcionais via query string:
    /api/produtos?categoria=Eletronicos&busca=fone
    """
    with lock:
        produtos = ler_produtos()

    categoria = request.args.get("categoria")
    busca = request.args.get("busca", "").lower()

    if categoria:
        produtos = [p for p in produtos if p["categoria"].lower() == categoria.lower()]
    if busca:
        produtos = [p for p in produtos if busca in p["nome"].lower()]

    return jsonify(produtos)


@app.route("/api/produtos/<int:produto_id>", methods=["GET"])
def obter_produto(produto_id):
    with lock:
        produtos = ler_produtos()
    produto = next((p for p in produtos if p["id"] == produto_id), None)
    if produto is None:
        abort(404, description="Produto nao encontrado")
    return jsonify(produto)


@app.route("/api/produtos", methods=["POST"])
def criar_produto():
    """Cria um novo produto. Corpo JSON esperado:
    { "nome": "...", "categoria": "...", "preco": 10.5, "estoque": 5,
      "descricao": "...", "imagem_url": "..." }
    """
    dados = request.get_json(force=True) or {}

    campos_obrigatorios = ["nome", "categoria", "preco", "estoque"]
    faltando = [c for c in campos_obrigatorios if c not in dados]
    if faltando:
        abort(400, description=f"Campos obrigatorios faltando: {', '.join(faltando)}")

    with lock:
        produtos = ler_produtos()
        novo = {
            "id": proximo_id(produtos),
            "nome": dados["nome"],
            "categoria": dados["categoria"],
            "preco": float(dados["preco"]),
            "estoque": int(dados["estoque"]),
            "descricao": dados.get("descricao", ""),
            "imagem_url": dados.get("imagem_url", ""),
        }
        produtos.append(novo)
        salvar_produtos(produtos)

    return jsonify(novo), 201


@app.route("/api/produtos/<int:produto_id>", methods=["PUT"])
def atualizar_produto(produto_id):
    dados = request.get_json(force=True) or {}

    with lock:
        produtos = ler_produtos()
        produto = next((p for p in produtos if p["id"] == produto_id), None)
        if produto is None:
            abort(404, description="Produto nao encontrado")

        for campo in ["nome", "categoria", "preco", "estoque", "descricao", "imagem_url"]:
            if campo in dados:
                produto[campo] = dados[campo]

        salvar_produtos(produtos)

    return jsonify(produto)


@app.route("/api/produtos/<int:produto_id>", methods=["DELETE"])
def remover_produto(produto_id):
    with lock:
        produtos = ler_produtos()
        existe = any(p["id"] == produto_id for p in produtos)
        if not existe:
            abort(404, description="Produto nao encontrado")

        produtos = [p for p in produtos if p["id"] != produto_id]
        salvar_produtos(produtos)

    return "", 204


@app.route("/api/produtos/exportar", methods=["GET"])
def exportar_planilha():
    """Permite baixar a planilha Excel atualizada diretamente."""
    with lock:
        garantir_planilha()
        caminho = EXCEL_PATH
    return send_file(
        caminho,
        as_attachment=True,
        download_name="produtos.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/categorias", methods=["GET"])
def listar_categorias():
    with lock:
        produtos = ler_produtos()
    categorias = sorted({p["categoria"] for p in produtos})
    return jsonify(categorias)


@app.route("/", methods=["GET"])
def index():
    if "application/json" in request.headers.get("Accept", "") and "text/html" not in request.headers.get("Accept", ""):
        return jsonify({
            "mensagem": "API do site de compras rodando.",
            "rotas": {
                "GET /api/produtos": "lista todos os produtos (filtros: ?categoria=&busca=)",
                "GET /api/produtos/<id>": "detalhe de um produto",
                "POST /api/produtos": "cria um novo produto",
                "PUT /api/produtos/<id>": "atualiza um produto",
                "DELETE /api/produtos/<id>": "remove um produto",
                "GET /api/produtos/exportar": "baixa a planilha Excel atual",
                "GET /api/categorias": "lista as categorias existentes",
            }
        })
    index_path = os.path.join(BASE_DIR, "index.html")
    if os.path.exists(index_path):
        return send_file(index_path)
    return jsonify({"mensagem": "API rodando."})


@app.route("/<path:filename>", methods=["GET"])
def arquivos_estaticos(filename):
    arquivo = os.path.join(BASE_DIR, filename)
    if os.path.exists(arquivo) and not os.path.isdir(arquivo):
        return send_from_directory(BASE_DIR, filename)
    abort(404)


if __name__ == "__main__":
    garantir_planilha()
    app.run(debug=True, host="0.0.0.0", port=5000)

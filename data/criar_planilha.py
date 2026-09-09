"""
Gera a planilha produtos.xlsx com dados de exemplo.
Execute apenas uma vez para criar o arquivo inicial (ou para resetar os dados).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "produtos"

headers = ["id", "nome", "categoria", "preco", "estoque", "descricao", "imagem_url"]

produtos = [
    (1, "Camiseta Basica Branca", "Roupas", 49.90, 120, "Camiseta 100% algodao, corte unissex", "https://picsum.photos/seed/1/400"),
    (2, "Tenis Esportivo Runner", "Calcados", 249.90, 45, "Tenis leve para corrida, amortecimento em gel", "https://picsum.photos/seed/2/400"),
    (3, "Fone de Ouvido Bluetooth", "Eletronicos", 189.90, 60, "Fone sem fio com cancelamento de ruido", "https://picsum.photos/seed/3/400"),
    (4, "Mochila Notebook 15pol", "Acessorios", 129.90, 30, "Mochila impermeavel com compartimento acolchoado", "https://picsum.photos/seed/4/400"),
    (5, "Garrafa Termica 1L", "Casa", 59.90, 80, "Mantem a temperatura por ate 12 horas", "https://picsum.photos/seed/5/400"),
    (6, "Relogio Smartwatch", "Eletronicos", 399.90, 25, "Monitor de saude e notificacoes no pulso", "https://picsum.photos/seed/6/400"),
    (7, "Jaqueta Corta Vento", "Roupas", 159.90, 40, "Jaqueta leve e resistente a agua", "https://picsum.photos/seed/7/400"),
    (8, "Cadeira Gamer", "Moveis", 899.90, 12, "Cadeira ergonomica com apoio lombar", "https://picsum.photos/seed/8/400"),
]

# Cabecalho
header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

for col, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Dados
body_font = Font(name="Arial", size=11)
for row_idx, produto in enumerate(produtos, start=2):
    for col_idx, value in enumerate(produto, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = body_font
        if headers[col_idx - 1] == "preco":
            cell.number_format = '"R$" #,##0.00'

# Largura das colunas
larguras = [6, 30, 15, 12, 10, 45, 35]
for i, largura in enumerate(larguras, start=1):
    ws.column_dimensions[get_column_letter(i)].width = largura

ws.freeze_panes = "A2"

wb.save("produtos.xlsx")
print("Planilha 'produtos.xlsx' criada com sucesso.")
